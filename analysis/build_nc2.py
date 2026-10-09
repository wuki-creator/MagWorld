import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(11)

lines = open('manuscript_nc.md', encoding='utf-8').read().split('\n')
figmap = {'Figure 1.': 'fig1_overview.png', 'Figure 2.': 'fig2_loco.png',
          'Figure 3.': 'fig3_decomposition.png', 'Figure 4.': 'fig4_dose.png',
          'Figure 5.': 'fig5_norman.png', 'Figure 6.': 'fig6_ablation.png',
          'Figure 7.': 'fig7_crest.png'}
for ln in lines:
    s = ln.strip()
    if not s:
        continue
    if s.startswith('# '):
        p = doc.add_heading(s[2:], level=0)
        continue
    if s.startswith('## '):
        doc.add_heading(s[3:], level=1); continue
    if s.startswith('### '):
        doc.add_heading(s[4:], level=2); continue
    img = None
    for k, v in figmap.items():
        if s.startswith(k):
            img = v; break
    if img:
        try:
            doc.add_picture(img, width=Inches(6.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            print('img fail', img, e)
        p = doc.add_paragraph(s)
        p.runs[0].italic = True
        continue
    doc.add_paragraph(s)
doc.save('VirtualCell_NC_full.docx')
print('docx saved')

# 投稿信
cl = Document()
cst = cl.styles['Normal']; cst.font.name = 'Times New Roman'; cst.font.size = Pt(11)
cover = """Dear Editors,

We are pleased to submit our manuscript entitled "VirtualCell: context-calibrated perturbation prediction exposes a background-residual structure of cellular responses" for consideration as an Article in Nature Communications.

Predicting cellular responses to perturbations is a central challenge toward computational "virtual cell" models. Our manuscript makes three contributions that we believe are of broad interest to the journal's readership. First, we identify and quantify a simple but previously uncharacterized statistical structure of perturbation data: roughly 62% of a perturbation signature's energy is a cell-line-specific background shared across perturbations, which explains why naive cross-context prediction fails (r = 0.14) and motivates a few-shot calibration strategy that recovers accuracy to r = 0.78, near the empirical noise ceiling (r = 0.85). Second, we introduce an evaluation protocol anchored to replicate-derived noise ceilings, showing that these ceilings vary three-fold with aggregation choices; we argue this should become standard practice in perturbation-prediction studies. Third, applying the framework's additive composition principle to the Norman et al. combinatorial CRISPR dataset, we show that double-perturbation outcomes are predictable at the noise ceiling (r = 0.91 vs. ceiling 0.93) and that additivity holds in expression space but not in latent space - a cautionary result for latent-arithmetic approaches.

All analyses use publicly available, harmonized datasets (sci-Plex, GEO: GSE225775; Norman et al. via scPerturb, Zenodo record 10044268), and we report negative results and ablations alongside positive findings. We confirm that this manuscript is original, has not been published previously, and is not under consideration elsewhere. The author declares no competing interests.

Thank you for your consideration. We look forward to your response.

Sincerely,
Wuzizhuo
Institution of HuahaiLi, AI Lab
"""
for para in cover.split('\n\n'):
    cl.add_paragraph(para.strip())
cl.save('CoverLetter_NC.docx')
print('cover letter saved')
